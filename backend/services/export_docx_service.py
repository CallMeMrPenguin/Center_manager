import os
from datetime import datetime
from typing import Dict, Any, List, Optional
import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from config.settings import get_setting
from database.db_manager import get_classes, get_class_attendance_grades
from database.utils import trunc_1_dec

def export_class_docx(
    class_id: int,
    date_str: Optional[str] = None,
    records: Optional[List[Dict[str, Any]]] = None,
    format_check_content_fn=None,
    get_session_test_config_fn=None,
    clean_num_fn=None,
) -> Dict[str, Any]:
    if not date_str:
        date_str = datetime.now().strftime("%Y-%m-%d")

    cls_list = get_classes()
    cls_info = next((c for c in cls_list if c["id"] == class_id), None)
    class_name = cls_info["class_name"] if cls_info else f"Class_{class_id}"

    attendance = records
    if not attendance:
        attendance = get_class_attendance_grades(class_id, date_str)

    session_cfg = get_session_test_config_fn(class_id, date_str) if get_session_test_config_fn else None
    c1_content = format_check_content_fn(session_cfg.get("check_1") if session_cfg else None) if format_check_content_fn else ""
    c2_content = format_check_content_fn(session_cfg.get("check_2") if session_cfg else None) if format_check_content_fn else ""

    doc = docx.Document()
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_p.add_run(f"BÁO CÁO NGHỈ HỌC & ĐIỂM BÀI HỌC\nLỚP: {class_name.upper()} - NGÀY: {date_str}")
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(30, 27, 75)

    info_parts = []
    if c1_content:
        info_parts.append(f"Check 1: {c1_content}")
    if c2_content:
        info_parts.append(f"Check 2: {c2_content}")
    if info_parts:
        cfg_p = doc.add_paragraph()
        cfg_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cfg_run = cfg_p.add_run("Nội dung kiểm tra: " + "   —   ".join(info_parts))
        cfg_run.bold = True
        cfg_run.font.size = Pt(11)
        cfg_run.font.color.rgb = RGBColor(49, 46, 129)

    doc.add_paragraph()
    c1_hdr = f"Check 1\n({c1_content})" if c1_content else "Check 1"
    c2_hdr = f"Check 2\n({c2_content})" if c2_content else "Check 2"
    headers = ["STT", "Họ và Tên", "Điểm Danh", c1_hdr, c2_hdr, "BTVN 1", "BTVN 2", "Luyện Đề"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        hdr_cells[i].paragraphs[0].runs[0].font.bold = True

    def _num(v):
        if clean_num_fn:
            return clean_num_fn(v)
        try:
            return float(v) if v is not None else 0.0
        except Exception:
            return 0.0

    for idx, r in enumerate(attendance, 1):
        c1 = _num(r.get("check_1"))
        c2 = _num(r.get("check_2"))
        hw1 = _num(r.get("homework"))
        hw2 = _num(r.get("homework_2"))
        mt = _num(r.get("mock_test"))

        row_cells = table.add_row().cells
        row_cells[0].text = str(idx)
        row_cells[1].text = str(r.get("student_name", ""))
        row_cells[2].text = str(r.get("status", "Có mặt"))
        row_cells[3].text = str(c1) if c1 > 0 else "-"
        row_cells[4].text = str(c2) if c2 > 0 else "-"
        row_cells[5].text = str(hw1) if hw1 > 0 else "-"
        row_cells[6].text = str(hw2) if hw2 > 0 else "-"
        row_cells[7].text = str(mt) if mt > 0 else "-"

    files_dir = get_setting("files_dir")
    if not files_dir or not os.path.exists(files_dir):
        files_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(__file__))), "workspace_files")
    os.makedirs(files_dir, exist_ok=True)
    ts = datetime.now().strftime("%H%M%S")
    filename = f"ClassReport_{class_name}_{date_str}_{ts}.docx"
    filepath = os.path.join(files_dir, filename)
    doc.save(filepath)
    return {"status": "success", "filename": filename, "filepath": filepath}
