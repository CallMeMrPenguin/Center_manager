import os
from typing import List, Dict, Any, Optional
from docx.shared import Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_TAB_ALIGNMENT, WD_COLOR_INDEX
from .common import cm_to_pt, set_cell_border

def docx_add_page_numbering(run):
    fldSimple = OxmlElement('w:fldSimple')
    fldSimple.set(qn('w:instr'), 'PAGE')
    run._r.append(fldSimple)

def docx_add_word_box(doc, words, layout):
    if not words:
        return
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = int(cm_to_pt(16.5) * 12700)
    cell = table.cell(0, 0)
    set_cell_border(cell, top="single", bottom="single", left="single", right="single", sz="8", val="single", color="808080")
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run("   ".join(words))
    run.font.name = layout.get("font_name", "Times New Roman")
    run.font.size = Pt(layout.get("font_size", 12.0) - 0.5)
    run.font.italic = True
    doc.add_paragraph()

def docx_add_test_header(doc, meta, layout):
    title = meta.get("title", "")
    school = meta.get("school", "")
    time_str = meta.get("time", "")
    exam_code = meta.get("exam_code", "")
    
    table = doc.add_table(rows=2, cols=2)
    table.autofit = False
    col_w = int(cm_to_pt(8.25) * 12700)
    for row in table.rows:
        for cell in row.cells:
            cell.width = col_w
            set_cell_border(cell, top="none", bottom="none", left="none", right="none")
            
    # Cell 0,0: School
    p00 = table.cell(0, 0).paragraphs[0]
    p00.paragraph_format.space_after = Pt(2)
    r_sch = p00.add_run(school.upper() if school else "TRƯỜNG THCS/THPT")
    r_sch.font.name = layout.get("font_name", "Times New Roman")
    r_sch.font.size = Pt(layout.get("font_size", 12.0))
    r_sch.font.bold = True
    
    # Cell 0,1: Title
    p01 = table.cell(0, 1).paragraphs[0]
    p01.paragraph_format.space_after = Pt(2)
    r_tit = p01.add_run(title.upper() if title else "BÀI KIỂM TRA ĐỊNH KỲ")
    r_tit.font.name = layout.get("font_name", "Times New Roman")
    r_tit.font.size = Pt(layout.get("font_size", 12.0))
    r_tit.font.bold = True
    
    # Cell 1,0: Time
    p10 = table.cell(1, 0).paragraphs[0]
    p10.paragraph_format.space_after = Pt(6)
    r_tim = p10.add_run(f"Thời gian: {time_str}" if time_str else "Thời gian: 45 phút")
    r_tim.font.name = layout.get("font_name", "Times New Roman")
    r_tim.font.size = Pt(layout.get("font_size", 12.0) - 1.0)
    r_tim.font.italic = True
    
    # Cell 1,1: Exam Code
    p11 = table.cell(1, 1).paragraphs[0]
    p11.paragraph_format.space_after = Pt(6)
    r_cod = p11.add_run(f"Mã đề: {exam_code}" if exam_code else "")
    r_cod.font.name = layout.get("font_name", "Times New Roman")
    r_cod.font.size = Pt(layout.get("font_size", 12.0) - 1.0)
    r_cod.font.bold = True

    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_after = Pt(8)

def docx_collect_answers(exercises: List[Dict[str, Any]]) -> List[Dict[str, str]]:
    answers = []
    for ex in exercises:
        q_num = str(ex.get("q", "")).strip()
        ans = ex.get("a", "")
        if isinstance(ans, list):
            ans = ", ".join(str(x) for x in ans)
        if q_num and ans:
            answers.append({"q": q_num, "a": str(ans)})
        if "k" in ex and isinstance(ex["k"], list):
            for sub in ex["k"]:
                sq = str(sub.get("q", "")).strip()
                sa = sub.get("a", "")
                if isinstance(sa, list):
                    sa = ", ".join(str(x) for x in sa)
                if sq and sa:
                    answers.append({"q": sq, "a": str(sa)})
    return answers

def docx_add_answer_key(doc, answers: List[Dict[str, str]], layout: Dict[str, Any]):
    if not answers:
        return
    doc.add_page_break()
    p_head = doc.add_paragraph()
    p_head.paragraph_format.space_before = Pt(14)
    p_head.paragraph_format.space_after = Pt(8)
    r_head = p_head.add_run("ĐÁP ÁN & HƯỚNG DẪN CHẤM")
    r_head.font.name = layout.get("font_name", "Times New Roman")
    r_head.font.size = Pt(layout.get("font_size", 12.0) + 1.0)
    r_head.font.bold = True

    num_cols = 5
    num_rows = (len(answers) + num_cols - 1) // num_cols
    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.autofit = True
    for idx, item in enumerate(answers):
        r_idx = idx % num_rows
        c_idx = idx // num_rows
        cell = table.cell(r_idx, c_idx)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(f"{item['q']}. {item['a']}")
        r.font.name = layout.get("font_name", "Times New Roman")
        r.font.size = Pt(layout.get("font_size", 12.0) - 1.0)

