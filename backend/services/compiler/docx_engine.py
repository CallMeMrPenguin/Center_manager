import os
import sys
import json
import math
import copy
import time
import random
from typing import List, Dict, Any, Optional
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_TAB_ALIGNMENT, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from .common import (
    INSTRUCTION_MAP_MCQ, INSTRUCTION_MAP_NON_MCQ, cm_to_pt,
    get_instruction_map, get_instruction_text, normalize_input_exercises,
    parse_text_formatting, clean_option_text, set_cell_border, get_unit_name
)
from .docx_layout import docx_add_page_numbering, docx_add_word_box, docx_add_test_header, docx_collect_answers, docx_add_answer_key

class WordDocumentCompilerDocx:
    """Fallback XML document compiler using python-docx."""
    def __init__(self, settings: Dict[str, Any] = None):
        self.settings = settings or {}
        self.doc = Document()
        self.is_answer_key = False
        self._configure_page()
        self._configure_styles()
        self._add_page_numbering()

    def _add_page_numbering(self):
        font_name = self.settings.get("font_name", "Times New Roman")
        font_size = self.settings.get("font_size", 12.0)
        
        for section in self.doc.sections:
            footer = section.footer
            p = footer.paragraphs[0]
            p.alignment = 1 # Center
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(0)
            p.text = ""
            
            def add_field(p_elem, field_name):
                fldSimple = OxmlElement('w:fldSimple')
                fldSimple.set(qn('w:instr'), field_name)
                
                r = OxmlElement('w:r')
                rPr = OxmlElement('w:rPr')
                
                rFonts = OxmlElement('w:rFonts')
                rFonts.set(qn('w:ascii'), font_name)
                rFonts.set(qn('w:hAnsi'), font_name)
                rPr.append(rFonts)
                
                sz = OxmlElement('w:sz')
                sz.set(qn('w:val'), str(int(font_size * 2)))
                rPr.append(sz)
                
                r.append(rPr)
                
                t = OxmlElement('w:t')
                t.text = "1"
                r.append(t)
                
                fldSimple.append(r)
                p_elem._p.append(fldSimple)
                
            add_field(p, 'PAGE')
            r_sep = p.add_run("/")
            r_sep.font.name = font_name
            r_sep.font.size = Pt(font_size)
            add_field(p, 'NUMPAGES')

    def _configure_page(self):
        top_cm = self.settings.get("margin_top", 2.0)
        bottom_cm = self.settings.get("margin_bottom", 2.0)
        left_cm = self.settings.get("margin_left", 3.0)
        right_cm = self.settings.get("margin_right", 1.5)
        
        for section in self.doc.sections:
            section.page_width = Cm(21)
            section.page_height = Cm(29.7)
            section.top_margin = Cm(top_cm)
            section.bottom_margin = Cm(bottom_cm)
            section.left_margin = Cm(left_cm)
            section.right_margin = Cm(right_cm)

    def _configure_styles(self):
        font_name = self.settings.get("font_name", "Times New Roman")
        font_size = self.settings.get("font_size", 12.0)
        line_spacing = self.settings.get("line_spacing", 1.15)
        space_after = self.settings.get("space_after", 6.0)
        
        style = self.doc.styles['Normal']
        font = style.font
        font.name = font_name
        font.size = Pt(font_size)
        
        p_format = style.paragraph_format
        p_format.space_before = Pt(0)
        p_format.space_after = Pt(space_after)
        p_format.line_spacing = line_spacing

    def add_instruction_header(self, text: str):
        space_before = self.settings.get("header_space_before", 14.0)
        space_after = self.settings.get("header_space_after", 8.0)
        font_size = self.settings.get("font_size", 12.0)
        
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.keep_with_next = True
        
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(font_size)

    def add_question(self, q_num: Any, q_text: str):
        space_before = self.settings.get("question_space_before", 6.0)
        space_after = self.settings.get("question_space_after", 4.0)
        
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.keep_with_next = True
        
        run_q = p.add_run(f"Question {q_num}: ")
        run_q.bold = True
        
        segments = parse_text_formatting(q_text)
        for seg in segments:
            run = p.add_run(seg["text"])
            if seg["bold"]: run.bold = True
            if seg["italic"]: run.italic = True
            if seg["underline"]: run.underline = True

    def add_options_grid(self, options: List[str], exercise_type: str, correct_ans: Any = None):
        if not options:
            return
        options = [clean_option_text(opt) for opt in options]
            
        correct_idx = -1
        if correct_ans:
            c_str = str(correct_ans).strip().upper()
            if len(c_str) == 1 and 'A' <= c_str <= 'E':
                correct_idx = ord(c_str) - ord('A')
            elif c_str in ['1', '2', '3', '4', '5']:
                correct_idx = int(c_str) - 1

        max_len = max(len(str(opt)) for opt in options) if options else 0
        if exercise_type in ["sg", "nt"]:
            cols = 1
        elif exercise_type in ["pr", "st"] or max_len < 15:
            cols = 4
        elif max_len < 35:
            cols = 2
        else:
            cols = 1

        left_indent_cm = self.settings.get("options_left_indent", 0.5)
        space_before = self.settings.get("options_space_before", 0.0)
        space_after = self.settings.get("options_space_after", 3.0)
        
        left_margin_cm = self.settings.get("margin_left", 3.0)
        right_margin_cm = self.settings.get("margin_right", 1.5)
        printable_width_cm = 21.0 - left_margin_cm - right_margin_cm
        remaining_width_cm = printable_width_cm - left_indent_cm

        if cols == 1:
            for idx, opt in enumerate(options):
                prefix = chr(65 + idx)
                p = self.doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(left_indent_cm)
                p.paragraph_format.space_before = Pt(space_before)
                p.paragraph_format.space_after = Pt(space_after)
                
                run_prefix = p.add_run(f"{prefix}. ")
                run_prefix.bold = True
                if idx == correct_idx and self.is_answer_key:
                    run_prefix.font.color.rgb = RGBColor(255, 0, 0)
                    run_prefix.font.highlight_color = WD_COLOR_INDEX.YELLOW
                
                segments = parse_text_formatting(opt)
                for seg in segments:
                    run = p.add_run(seg["text"])
                    if seg["bold"]: run.bold = True
                    if seg["italic"]: run.italic = True
                    if seg["underline"]: run.underline = True
                    if idx == correct_idx and self.is_answer_key:
                        run.font.color.rgb = RGBColor(255, 0, 0)
                        run.font.highlight_color = WD_COLOR_INDEX.YELLOW

        elif cols == 2:
            col_width = remaining_width_cm / 2
            tab_pos = left_indent_cm + col_width
            
            for row in range(2):
                p = self.doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(left_indent_cm)
                p.paragraph_format.space_before = Pt(space_before)
                p.paragraph_format.space_after = Pt(space_after)
                p.paragraph_format.tab_stops.add_tab_stop(Cm(tab_pos), WD_TAB_ALIGNMENT.LEFT)
                
                idx1 = row * 2
                if idx1 < len(options):
                    prefix = chr(65 + idx1)
                    run_prefix = p.add_run(f"{prefix}. ")
                    run_prefix.bold = True
                    if idx1 == correct_idx and self.is_answer_key:
                        run_prefix.font.color.rgb = RGBColor(255, 0, 0)
                        run_prefix.font.highlight_color = WD_COLOR_INDEX.YELLOW
                    for seg in parse_text_formatting(options[idx1]):
                        run = p.add_run(seg["text"])
                        if seg["bold"]: run.bold = True
                        if seg["italic"]: run.italic = True
                        if seg["underline"]: run.underline = True
                        if idx1 == correct_idx and self.is_answer_key:
                            run.font.color.rgb = RGBColor(255, 0, 0)
                            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
                        
                r_tab = p.add_run("\t")
                r_tab.font.underline = False
                
                idx2 = row * 2 + 1
                if idx2 < len(options):
                    prefix = chr(65 + idx2)
                    run_prefix = p.add_run(f"{prefix}. ")
                    run_prefix.bold = True
                    if idx2 == correct_idx and self.is_answer_key:
                        run_prefix.font.color.rgb = RGBColor(255, 0, 0)
                        run_prefix.font.highlight_color = WD_COLOR_INDEX.YELLOW
                    for seg in parse_text_formatting(options[idx2]):
                        run = p.add_run(seg["text"])
                        if seg["bold"]: run.bold = True
                        if seg["italic"]: run.italic = True
                        if seg["underline"]: run.underline = True
                        if idx2 == correct_idx and self.is_answer_key:
                            run.font.color.rgb = RGBColor(255, 0, 0)
                            run.font.highlight_color = WD_COLOR_INDEX.YELLOW

        elif cols == 4:
            col_width = remaining_width_cm / 4
            tab1 = left_indent_cm + col_width
            tab2 = left_indent_cm + (col_width * 2)
            tab3 = left_indent_cm + (col_width * 3)
            
            p = self.doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(left_indent_cm)
            p.paragraph_format.space_before = Pt(space_before)
            p.paragraph_format.space_after = Pt(space_after)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(tab1), WD_TAB_ALIGNMENT.LEFT)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(tab2), WD_TAB_ALIGNMENT.LEFT)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(tab3), WD_TAB_ALIGNMENT.LEFT)
            
            for idx in range(4):
                if idx < len(options):
                    prefix = chr(65 + idx)
                    run_prefix = p.add_run(f"{prefix}. ")
                    run_prefix.bold = True
                    if idx == correct_idx and self.is_answer_key:
                        run_prefix.font.color.rgb = RGBColor(255, 0, 0)
                        run_prefix.font.highlight_color = WD_COLOR_INDEX.YELLOW
                    for seg in parse_text_formatting(options[idx]):
                        run = p.add_run(seg["text"])
                        if seg["bold"]: run.bold = True
                        if seg["italic"]: run.italic = True
                        if seg["underline"]: run.underline = True
                        if idx == correct_idx and self.is_answer_key:
                            run.font.color.rgb = RGBColor(255, 0, 0)
                            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
                
                if idx < 3:
                    r_tab = p.add_run("\t")
                    r_tab.font.underline = False
            
    def add_word_box(self, words: List[str]):
        return docx_add_word_box(self.doc, words, self.layout)

    def add_test_header(self, meta: Dict[str, Any]):
        return docx_add_test_header(self.doc, meta, self.layout)

    def collect_answers(self, exercises: List[Dict[str, Any]]) -> List[Dict[str, str]]:
        return docx_collect_answers(exercises)

    def add_answer_key(self, answers: List[Dict[str, str]]):
        return docx_add_answer_key(self.doc, answers, self.layout)

    def compile_exercises(self, exercises: List[Dict[str, Any]], grade: str = "", unit: str = "", version_code: str = "", include_answer_key: bool = True):
        exercises = normalize_input_exercises(exercises)

        if grade or unit or version_code:
            self.add_test_header(grade, unit, version_code)

        current_block_key = None
        for ex in exercises:
            ex_type = ex.get("t", "mq")
            options = ex.get("o", [])
            if not isinstance(options, list):
                options = []
            is_mcq = len(options) >= 2
            
            block_key = (ex_type, is_mcq)
            if block_key != current_block_key:
                current_block_key = block_key
                instruction_text = get_instruction_text(ex_type, is_mcq)
                if instruction_text:
                    self.add_instruction_header(instruction_text)
            
            if ex_type == "wb":
                words = ex.get("w", [])
                if words:
                    self.add_word_box(words)
                space_before = self.settings.get("passage_space_before", 4.0)
                space_after = self.settings.get("passage_space_after", 6.0)
                indent_first = self.settings.get("passage_indent_first", 0.75)
                
                for paragraph_text in ex.get("b", []):
                    p_passage = self.doc.add_paragraph()
                    p_passage.paragraph_format.space_before = Pt(space_before)
                    p_passage.paragraph_format.space_after = Pt(space_after)
                    p_passage.paragraph_format.first_line_indent = Cm(indent_first)
                    
                    segments = parse_text_formatting(paragraph_text)
                    for seg in segments:
                        run = p_passage.add_run(seg["text"])
                        if seg["bold"]: run.bold = True
                        if seg["italic"]: run.italic = True
                        if seg["underline"]: run.underline = True
                
                for sub in ex.get("k", []):
                    sub_q = sub.get("q")
                    sub_x = sub.get("x", "")
                    sub_o = sub.get("o", [])
                    sub_a = sub.get("a", "")
                    if sub_q and (sub_x.strip() != "" or sub_o):
                        self.add_question(sub_q, sub_x)
                        if sub_o:
                            self.add_options_grid(sub_o, ex_type, correct_ans=sub_a)

            elif ex_type in ["cz", "rd"]:
                space_before = self.settings.get("passage_space_before", 4.0)
                space_after = self.settings.get("passage_space_after", 6.0)
                indent_first = self.settings.get("passage_indent_first", 0.75)
                
                for paragraph_text in ex.get("b", []):
                    p_passage = self.doc.add_paragraph()
                    p_passage.paragraph_format.space_before = Pt(space_before)
                    p_passage.paragraph_format.space_after = Pt(space_after)
                    p_passage.paragraph_format.first_line_indent = Cm(indent_first)
                    
                    segments = parse_text_formatting(paragraph_text)
                    for seg in segments:
                        run = p_passage.add_run(seg["text"])
                        if seg["bold"]: run.bold = True
                        if seg["italic"]: run.italic = True
                        if seg["underline"]: run.underline = True
                
                for sub in ex.get("k", []):
                    sub_q = sub.get("q")
                    sub_x = sub.get("x", "")
                    sub_o = sub.get("o", [])
                    sub_a = sub.get("a", "")
                    if sub_q and (sub_x.strip() != "" or sub_o):
                        self.add_question(sub_q, sub_x)
                        if sub_o:
                            self.add_options_grid(sub_o, ex_type, correct_ans=sub_a)
                    
            elif ex_type == "ro":
                q_num = ex.get("q")
                q_text = ex.get("x", "Choose the best arrangement of the sentences:")
                self.add_question(q_num, q_text)
                
                space_before = self.settings.get("reorder_space_before", 0.0)
                space_after = self.settings.get("reorder_space_after", 2.0)
                left_indent = self.settings.get("reorder_left_indent", 1.0)
                
                for item in ex.get("i", []):
                    p_item = self.doc.add_paragraph()
                    p_item.paragraph_format.left_indent = Cm(left_indent)
                    p_item.paragraph_format.space_before = Pt(space_before)
                    p_item.paragraph_format.space_after = Pt(space_after)
                    
                    segments = parse_text_formatting(item)
                    for seg in segments:
                        run = p_item.add_run(seg["text"])
                        if seg["bold"]: run.bold = True
                        if seg["italic"]: run.italic = True
                        if seg["underline"]: run.underline = True
                
                self.add_options_grid(ex.get("o", []), ex_type, correct_ans=ex.get("a", ""))
                
            elif ex_type == "nt":
                q_num = ex.get("q")
                q_text = ex.get("x", "")
                self.doc.add_paragraph()
                self.add_question(q_num, q_text)
                
                space_before = self.settings.get("notice_space_before", 4.0)
                space_after = self.settings.get("notice_space_after", 6.0)
                left_indent = self.settings.get("notice_left_indent", 1.0)
                
                p_body = self.doc.add_paragraph()
                p_body.paragraph_format.left_indent = Cm(left_indent)
                p_body.paragraph_format.space_before = Pt(space_before)
                p_body.paragraph_format.space_after = Pt(space_after)
                
                segments = parse_text_formatting(ex.get("b", ""))
                for seg in segments:
                    run = p_body.add_run(seg["text"])
                    run.italic = True
                    if seg["bold"]: run.bold = True
                    if seg["underline"]: run.underline = True
                
                self.add_options_grid(ex.get("o", []), ex_type, correct_ans=ex.get("a", ""))
                
            else:
                q_num = ex.get("q")
                q_text = ex.get("x", "")
                options = ex.get("o", [])
                ans = ex.get("a", "")
                
                self.add_question(q_num, q_text)
                self.add_options_grid(options, ex_type, correct_ans=ans)

        if include_answer_key:
            self.add_answer_key(exercises)

    def compile(self, exercises: List[Dict[str, Any]], output_filepath: Any, grade: str = "", unit: str = "", version_code: str = "", include_answer_key: bool = True, is_answer_key: bool = False, is_test: bool = False):
        self.is_answer_key = is_answer_key
        if grade or unit or version_code or is_test:
            self.add_test_header(grade=grade, unit=unit, version_code=version_code, is_test=is_test)
        self.compile_exercises(exercises, grade=grade, unit=unit, version_code=version_code, include_answer_key=include_answer_key)
        self.doc.save(output_filepath)


