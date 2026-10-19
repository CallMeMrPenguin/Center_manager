import os
import re
from datetime import datetime
from typing import Dict, Any, List, Optional
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from config.settings import get_setting
from database.db_manager import get_classes, get_class_attendance_grades
from database.utils import trunc_1_dec

def export_class_excel(
    class_id: int,
    date_str: Optional[str] = None,
    records: Optional[List[Dict[str, Any]]] = None,
    thresholds: Optional[Dict[str, Any]] = None,
    test_config: Optional[Dict[str, Any]] = None,
    visible_columns: Optional[List[str]] = None,
    format_check_content_fn=None,
    get_session_test_config_fn=None,
    clean_num_fn=None,
) -> Dict[str, Any]:
    if not date_str:
        date_str = datetime.now().strftime("%Y-%m-%d")
    ts = datetime.now().strftime("%S-%M-%H -%d=%m-%y")

    cls_list = get_classes()
    cls_info = next((c for c in cls_list if c["id"] == class_id), None)
    class_name = cls_info["class_name"] if cls_info else f"Class_{class_id}"

    attendance = records if records is not None else get_class_attendance_grades(class_id, date_str)

    get_test_cfg = get_session_test_config_fn or (lambda cid, d: None)
    fmt_check = format_check_content_fn or (lambda cfg: "")
    c_num = clean_num_fn or (lambda v: float(v) if v else 0.0)

    session_cfg = test_config if test_config else get_test_cfg(class_id, date_str)
    c1_content = fmt_check(session_cfg.get("check_1") if session_cfg else None)
    c2_content = fmt_check(session_cfg.get("check_2") if session_cfg else None)

    # Pre-calculate class averages and thresholds
    c1_vals = [c_num(r.get("check_1")) for r in attendance if c_num(r.get("check_1")) > 0 and str(r.get("status")) != "Vắng mặt"]
    c2_vals = [c_num(r.get("check_2")) for r in attendance if c_num(r.get("check_2")) > 0 and str(r.get("status")) != "Vắng mặt"]
    hw1_vals = [c_num(r.get("homework")) for r in attendance if c_num(r.get("homework")) > 0 and str(r.get("status")) != "Vắng mặt"]
    hw2_vals = [c_num(r.get("homework_2")) for r in attendance if c_num(r.get("homework_2")) > 0 and str(r.get("status")) != "Vắng mặt"]
    mt_vals = [c_num(r.get("mock_test")) for r in attendance if c_num(r.get("mock_test")) > 0 and str(r.get("status")) != "Vắng mặt"]

    calc_avg_1 = trunc_1_dec(sum(c1_vals) / len(c1_vals)) if c1_vals else 0.0
    calc_avg_2 = trunc_1_dec(sum(c2_vals) / len(c2_vals)) if c2_vals else 0.0
    calc_avg_hw1 = trunc_1_dec(sum(hw1_vals) / len(hw1_vals)) if hw1_vals else 0.0
    calc_avg_hw2 = trunc_1_dec(sum(hw2_vals) / len(hw2_vals)) if hw2_vals else 0.0
    calc_avg_mt = trunc_1_dec(sum(mt_vals) / len(mt_vals)) if mt_vals else 0.0

    th = thresholds or {}
    def _pick_thresh(key: str, calc_val: float) -> float:
        if key in th and th[key] is not None and str(th[key]).strip() != '':
            try:
                v = float(th[key])
                if v > 0:
                    return trunc_1_dec(v)
            except (ValueError, TypeError):
                pass
        return calc_val

    t_c1 = _pick_thresh("check_1", calc_avg_1)
    t_c2 = _pick_thresh("check_2", calc_avg_2)
    t_hw1 = _pick_thresh("homework", calc_avg_hw1)
    t_hw2 = _pick_thresh("homework_2", calc_avg_hw2)
    t_mt = _pick_thresh("mock_test", calc_avg_mt)

    check_avgs = [a for a in (t_c1, t_c2) if a > 0]
    check_combined = sum(check_avgs) / len(check_avgs) if check_avgs else 0.0
    avg_diff = trunc_1_dec(abs(t_hw1 - check_combined)) if (t_hw1 > 0 and check_combined > 0) else 0.0

    # Styles
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    header_font = Font(name="Times New Roman", color="FFFFFF", bold=True, size=12)
    data_font = Font(name="Times New Roman", size=12)
    name_font = Font(name="Times New Roman", size=12, bold=True, color="0F172A")

    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    fill_8 = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    font_8 = Font(name="Times New Roman", size=12, color="15803D", bold=True)
    fill_65 = PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid")
    font_65 = Font(name="Times New Roman", size=12, color="0369A1", bold=True)
    fill_5 = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    font_5 = Font(name="Times New Roman", size=12, color="B45309", bold=True)
    fill_under5 = PatternFill(start_color="FFE4E6", end_color="FFE4E6", fill_type="solid")
    font_under5 = Font(name="Times New Roman", size=12, color="BE123C", bold=True)
    fill_abs = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    font_abs = Font(name="Times New Roman", size=12, color="64748B", bold=True)

    def _render_score_cell(cell, val, row_bg):
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center")
        if val > 0:
            cell.value = val
            cell.number_format = '0.0'
            if val >= 8.0:
                cell.fill, cell.font = fill_8, font_8
            elif val >= 6.5:
                cell.fill, cell.font = fill_65, font_65
            elif val >= 5.0:
                cell.fill, cell.font = fill_5, font_5
            else:
                cell.fill, cell.font = fill_under5, font_under5
        else:
            cell.value = "-"
            cell.font = data_font
            cell.fill = row_bg

    c1_hdr = f"Check 1\n({c1_content})" if c1_content else "Check 1"
    c2_hdr = f"Check 2\n({c2_content})" if c2_content else "Check 2"

    ALL_COLUMNS_CONFIG = {
        "stt": {
            "id": "stt",
            "header": "STT",
            "width": 8,
            "render_cell": lambda cell, ctx: _render_cell_val(cell, ctx["idx"], data_font, ctx["row_bg"], thin_border),
            "avg_value": "",
        },
        "student_name": {
            "id": "student_name",
            "header": "Họ và Tên",
            "width": 26,
            "render_cell": lambda cell, ctx: _render_cell_val(cell, ctx["st_name"], name_font, ctx["row_bg"], thin_border),
            "avg_value": "Điểm trung bình (Average)",
        },
        "status": {
            "id": "status",
            "header": "Điểm Danh",
            "width": 14,
            "render_cell": lambda cell, ctx: _render_cell_val(cell, ctx["status_val"], data_font, ctx["row_bg"], thin_border),
            "avg_value": "",
        },
        "check_1": {
            "id": "check_1",
            "header": c1_hdr,
            "width": 35 if c1_content else 14,
            "render_cell": lambda cell, ctx: _render_score_cell(cell, ctx["c1"], ctx["row_bg"]),
            "avg_value": t_c1 if t_c1 > 0 else "-",
        },
        "check_2": {
            "id": "check_2",
            "header": c2_hdr,
            "width": 35 if c2_content else 14,
            "render_cell": lambda cell, ctx: _render_score_cell(cell, ctx["c2"], ctx["row_bg"]),
            "avg_value": t_c2 if t_c2 > 0 else "-",
        },
        "homework": {
            "id": "homework",
            "header": "BTVN 1",
            "width": 13,
            "render_cell": lambda cell, ctx: _render_score_cell(cell, ctx["hw1"], ctx["row_bg"]),
            "avg_value": t_hw1 if t_hw1 > 0 else "-",
        },
        "homework_2": {
            "id": "homework_2",
            "header": "BTVN 2",
            "width": 13,
            "render_cell": lambda cell, ctx: _render_score_cell(cell, ctx["hw2"], ctx["row_bg"]),
            "avg_value": t_hw2 if t_hw2 > 0 else "-",
        },
        "mock_test": {
            "id": "mock_test",
            "header": "Luyện Đề",
            "width": 13,
            "render_cell": lambda cell, ctx: _render_score_cell(cell, ctx["mt"], ctx["row_bg"]),
            "avg_value": t_mt if t_mt > 0 else "-",
        },
        "score_discrepancy": {
            "id": "score_discrepancy",
            "header": "Độ Lệch",
            "width": 15,
            "render_cell": lambda cell, ctx: _render_diff_cell(cell, ctx["diff_val"], ctx["row_bg"], thin_border, data_font, fill_under5, font_under5, fill_5, font_5, fill_8, font_8),
            "avg_value": avg_diff if avg_diff > 0 else "-",
        },
        "need_improvement": {
            "id": "need_improvement",
            "header": "Cần Cố Gắng (Dưới TB)",
            "width": 48,
            "render_cell": lambda cell, ctx: _render_eval_cell(cell, ctx, thin_border, fill_abs, font_abs, fill_under5, font_under5, fill_8, font_8),
            "avg_value": "Đã tính TB lớp",
        },
    }

    COLUMN_ALIASES = {
        "stt": "stt",
        "student_name": "student_name",
        "full_name": "student_name",
        "name": "student_name",
        "status": "status",
        "attendance": "status",
        "check_1": "check_1",
        "vocab": "check_1",
        "check_2": "check_2",
        "grammar": "check_2",
        "homework": "homework",
        "homework_1": "homework",
        "btvn_1": "homework",
        "btvn1": "homework",
        "hw1": "homework",
        "homework_2": "homework_2",
        "btvn_2": "homework_2",
        "btvn2": "homework_2",
        "hw2": "homework_2",
        "mock_test": "mock_test",
        "luyen_de": "mock_test",
        "mock": "mock_test",
        "score_discrepancy": "score_discrepancy",
        "do_lech": "score_discrepancy",
        "divergence": "score_discrepancy",
        "diff": "score_discrepancy",
        "need_improvement": "need_improvement",
        "can_co_gang": "need_improvement",
        "evaluation": "need_improvement",
    }

    active_columns = []
    if visible_columns:
        for cid in visible_columns:
            clean_id = str(cid).lower().strip()
            canonical_id = COLUMN_ALIASES.get(clean_id)
            if canonical_id and canonical_id in ALL_COLUMNS_CONFIG:
                cfg = ALL_COLUMNS_CONFIG[canonical_id]
                if cfg not in active_columns:
                    active_columns.append(cfg)

    # Fallback to default full order if empty or not provided
    if not active_columns:
        default_order = [
            "stt", "student_name", "status", "check_1", "check_2",
            "homework", "homework_2", "mock_test", "score_discrepancy", "need_improvement"
        ]
        active_columns = [ALL_COLUMNS_CONFIG[k] for k in default_order if k in ALL_COLUMNS_CONFIG]

    num_cols = len(active_columns)
    last_col_letter = get_column_letter(num_cols)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "DiemDanh"

    # Row 1: Master Header
    if num_cols > 1:
        ws.merge_cells(f"A1:{last_col_letter}1")
    ws["A1"] = f"BÁO CÁO ĐIỂM DANH & ĐIỂM BÀI HỌC - {class_name.upper()} ({date_str})"
    ws["A1"].font = Font(name="Times New Roman", size=14, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill(start_color="1D4ED8", end_color="1D4ED8", fill_type="solid")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 38

    # Row 2: Topic subtitle strip
    info_parts = []
    if c1_content:
        info_parts.append(f"Check 1: {c1_content}")
    if c2_content:
        info_parts.append(f"Check 2: {c2_content}")

    if info_parts:
        if num_cols > 1:
            ws.merge_cells(f"A2:{last_col_letter}2")
        ws["A2"] = "Nội dung kiểm tra: " + "   —   ".join(info_parts)
        ws["A2"].font = Font(name="Times New Roman", size=11, bold=True, color="1E40AF")
        ws["A2"].fill = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")
        ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 26
    else:
        ws.cell(row=2, column=1, value="")

    # Row 3: Headers
    max_topic_len = max(len(c1_content or ""), len(c2_content or ""))
    ws.row_dimensions[3].height = 64 if max_topic_len > 35 else (52 if max_topic_len > 0 else 32)

    for col_idx, col_cfg in enumerate(active_columns, 1):
        cell = ws.cell(row=3, column=col_idx, value=col_cfg["header"])
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    start_row = 4
    end_row = start_row + len(attendance) - 1 if len(attendance) > 0 else start_row
    avg_row_idx = end_row + 1

    # Render student rows
    for idx, r in enumerate(attendance, 1):
        curr_row = start_row + idx - 1
        st_name = str(r.get("student_name", ""))
        status_val = str(r.get("status", "Có mặt"))
        c1 = c_num(r.get("check_1"))
        c2 = c_num(r.get("check_2"))
        hw1 = c_num(r.get("homework"))
        hw2 = c_num(r.get("homework_2"))
        mt = c_num(r.get("mock_test"))

        ws.row_dimensions[curr_row].height = 24
        row_bg = white_fill if (idx % 2 == 1) else zebra_fill

        diff_val = None
        if status_val != "Vắng mặt" and hw1 > 0:
            c_scores = [s for s in (c1, c2) if s > 0]
            if c_scores:
                c_avg = sum(c_scores) / len(c_scores)
                diff_val = trunc_1_dec(abs(hw1 - c_avg))

        below_subs = []
        if c1 > 0 and t_c1 > 0 and c1 < t_c1:
            below_subs.append("Check 1")
        if c2 > 0 and t_c2 > 0 and c2 < t_c2:
            below_subs.append("Check 2")
        if hw1 > 0 and t_hw1 > 0 and hw1 < t_hw1:
            below_subs.append("BTVN 1")
        if hw2 > 0 and t_hw2 > 0 and hw2 < t_hw2:
            below_subs.append("BTVN 2")
        if mt > 0 and t_mt > 0 and mt < t_mt:
            below_subs.append("Luyện Đề")

        ctx = {
            "idx": idx,
            "st_name": st_name,
            "status_val": status_val,
            "c1": c1,
            "c2": c2,
            "hw1": hw1,
            "hw2": hw2,
            "mt": mt,
            "diff_val": diff_val,
            "below_subs": below_subs,
            "row_bg": row_bg,
        }

        for col_idx, col_cfg in enumerate(active_columns, 1):
            cell = ws.cell(row=curr_row, column=col_idx)
            col_cfg["render_cell"](cell, ctx)

    # Auto Filter on Header (Row 3)
    if len(attendance) > 0:
        ws.auto_filter.ref = f"A3:{last_col_letter}{end_row}"

    # Average / Threshold Row
    avg_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    avg_font = Font(name="Times New Roman", bold=True, size=12, color="92400E")
    ws.row_dimensions[avg_row_idx].height = 26

    for col_idx, col_cfg in enumerate(active_columns, 1):
        val = col_cfg["avg_value"]
        cell = ws.cell(row=avg_row_idx, column=col_idx, value=val)
        if isinstance(val, (int, float)) and val > 0:
            cell.number_format = '0.0'
        cell.font = avg_font
        cell.fill = avg_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Spacing row
    ws.row_dimensions[avg_row_idx + 1].height = 12

    # Summary rows (Below Average Lists)
    candidate_labels = [
        ("Check 1", t_c1, [r.get("student_name") for r in attendance if c_num(r.get("check_1")) > 0 and c_num(r.get("check_1")) < t_c1 and str(r.get("status")) != "Vắng mặt"]),
        ("Check 2", t_c2, [r.get("student_name") for r in attendance if c_num(r.get("check_2")) > 0 and c_num(r.get("check_2")) < t_c2 and str(r.get("status")) != "Vắng mặt"]),
        ("BTVN", t_hw1, [r.get("student_name") for r in attendance if c_num(r.get("homework")) > 0 and c_num(r.get("homework")) < t_hw1 and str(r.get("status")) != "Vắng mặt"]),
    ]
    if t_hw2 > 0 or len(hw2_vals) > 0:
        candidate_labels.append(("BTVN 2", t_hw2, [r.get("student_name") for r in attendance if c_num(r.get("homework_2")) > 0 and c_num(r.get("homework_2")) < t_hw2 and str(r.get("status")) != "Vắng mặt"]))
    if t_mt > 0 or len(mt_vals) > 0:
        candidate_labels.append(("Luyện Đề", t_mt, [r.get("student_name") for r in attendance if c_num(r.get("mock_test")) > 0 and c_num(r.get("mock_test")) < t_mt and str(r.get("status")) != "Vắng mặt"]))

    for idx, (m_label, thresh_val, below_list) in enumerate(candidate_labels):
        r_idx = avg_row_idx + 2 + idx
        ws.row_dimensions[r_idx].height = 28
        title_text = f"{m_label} dưới TB (< {thresh_val})" if thresh_val > 0 else m_label
        names_text = ", ".join([str(n) for n in below_list if n]) if below_list else "Không có (Tất cả đạt)"

        if num_cols >= 3:
            ws.merge_cells(f"A{r_idx}:B{r_idx}")
            sum_title = ws.cell(row=r_idx, column=1, value=title_text)
            sum_title.font = Font(name="Times New Roman", bold=True, color="7F1D1D", size=12)
            sum_title.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            sum_title.border = thin_border
            ws.cell(row=r_idx, column=2).border = thin_border

            ws.merge_cells(f"C{r_idx}:{last_col_letter}{r_idx}")
            val_cell = ws.cell(row=r_idx, column=3, value=names_text)
            val_cell.font = Font(name="Times New Roman", bold=True, color="1E1E2F", size=12)
            val_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            for cn in range(3, num_cols + 1):
                ws.cell(row=r_idx, column=cn).border = thin_border
        elif num_cols == 2:
            ws.cell(row=r_idx, column=1, value=title_text).border = thin_border
            ws.cell(row=r_idx, column=2, value=names_text).border = thin_border
        else:
            ws.cell(row=r_idx, column=1, value=f"{title_text}: {names_text}").border = thin_border

    # Set column widths
    for col_idx, col_cfg in enumerate(active_columns, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = col_cfg["width"]

    files_dir = get_setting("files_dir")
    if not files_dir or not os.path.exists(files_dir):
        files_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(__file__))), "workspace_files")
    os.makedirs(files_dir, exist_ok=True)
    filename = f"ClassReport_{class_name}_{date_str}_{ts}.xlsx"
    filepath = os.path.join(files_dir, filename)
    wb.save(filepath)
    return {"status": "success", "filename": filename, "filepath": filepath}

def _render_cell_val(cell, val, font, fill, border):
    cell.value = val
    cell.font = font
    cell.fill = fill
    cell.border = border
    cell.alignment = Alignment(horizontal="center", vertical="center")

def _render_diff_cell(cell, diff_val, row_bg, border, data_font, fill_under5, font_under5, fill_5, font_5, fill_8, font_8):
    cell.border = border
    cell.alignment = Alignment(horizontal="center", vertical="center")
    if diff_val is not None:
        cell.value = diff_val
        cell.number_format = '0.0'
        if diff_val >= 2.0:
            cell.fill, cell.font = fill_under5, font_under5
        elif diff_val >= 1.0:
            cell.fill, cell.font = fill_5, font_5
        else:
            cell.fill, cell.font = fill_8, font_8
    else:
        cell.value = "-"
        cell.font = data_font
        cell.fill = row_bg

def _render_eval_cell(cell, ctx, border, fill_abs, font_abs, fill_under5, font_under5, fill_8, font_8):
    cell.border = border
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if ctx["status_val"] == "Vắng mặt":
        cell.value = "Vắng mặt"
        cell.fill, cell.font = fill_abs, font_abs
    else:
        below_subs = ctx["below_subs"]
        if below_subs:
            cell.value = f"Cần cố gắng ({', '.join(below_subs)})"
            cell.fill, cell.font = fill_under5, font_under5
        else:
            cell.value = "Đạt yêu cầu"
            cell.fill, cell.font = fill_8, font_8
